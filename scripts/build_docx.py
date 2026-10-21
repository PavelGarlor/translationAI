#!/usr/bin/env python3
"""Build a formatted .docx from a translation written in simple markdown.

Usage: python scripts/build_docx.py <translation.md> <output.docx>

Supported syntax (one block per line unless noted):
  # / ## / ###          headings (bold, centred for #)
  [center] text         centred paragraph
  [right] text          right-aligned paragraph
  | a | b |             table rows (consecutive lines form one table;
                        a |---| separator row is ignored)
  [pagebreak]           page break (mirror the original's pages)
  **bold**, *italic*    inline formatting
  blank line            paragraph separator; extra blank lines add spacing
Everything else is a normal paragraph. Bracketed descriptions such as
/Rundsiegel/ or /Unterschrift/ are kept as written.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt

FONT = "Times New Roman"
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def add_runs(par, text: str) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def setup(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(4)
    for s in doc.sections:
        s.page_height, s.page_width = Cm(29.7), Cm(21.0)
        s.left_margin = s.right_margin = Cm(2.5)
        s.top_margin = s.bottom_margin = Cm(2.0)


def flush_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c in range(ncols):
            cell = table.cell(r, c)
            cell.paragraphs[0].text = ""
            add_runs(cell.paragraphs[0], row[c] if c < len(row) else "")
    rows.clear()


def build(md: str, out: Path) -> None:
    doc = Document()
    setup(doc)
    table_rows: list[list[str]] = []
    blank_run = 0

    for raw in md.splitlines():
        line = raw.rstrip()
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells if c):
                table_rows.append(cells)
            continue
        flush_table(doc, table_rows)

        if not line.strip():
            blank_run += 1
            if blank_run > 1:
                doc.add_paragraph()
            continue
        blank_run = 0

        if line.strip() == "[pagebreak]":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            continue

        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            level = len(m.group(1))
            p = doc.add_paragraph()
            run = p.add_run(m.group(2))
            run.bold = True
            run.font.size = Pt({1: 14, 2: 13, 3: 12}[level])
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue

        align = None
        if line.startswith("[center]"):
            align, line = WD_ALIGN_PARAGRAPH.CENTER, line[8:].lstrip()
        elif line.startswith("[right]"):
            align, line = WD_ALIGN_PARAGRAPH.RIGHT, line[7:].lstrip()
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        add_runs(p, line)

    flush_table(doc, table_rows)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"Saved {out}")


def main() -> None:
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    build(Path(sys.argv[1]).read_text(encoding="utf-8"), Path(sys.argv[2]))


if __name__ == "__main__":
    main()
